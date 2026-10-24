"""Crear el Word de la Parte 2 usando los CSV de R y los PNG de Matplotlib.

Uso: python python/crear_informe.py --salida "carpeta de entrega"
No genera carpetas docs/datos ni archivos Markdown.
"""
from pathlib import Path
import argparse
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
T = ROOT / "resultados/tablas"
F = ROOT / "resultados/figuras"
parser = argparse.ArgumentParser()
parser.add_argument("--salida", type=Path, default=ROOT / "resultados")
args = parser.parse_args()
args.salida.mkdir(parents=True, exist_ok=True)
OUT = args.salida / "Proyecto_Integrador_Multivariada_PARTE2.docx"
def leer(n): return pd.read_csv(T / f"{n}.csv")
def f(x, d=2):
    if pd.isna(x): return "—"
    return f"{float(x):,.{d}f}".replace(",", "@").replace(".", ",").replace("@", ".")
def abreviar(x): return str(x).removeprefix("log_")

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(2.5)
sec.header_distance = sec.footer_distance = Cm(1.2)
for name in ["Normal", "Title", "Heading 1", "Heading 2", "Caption"]:
    style = doc.styles[name]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor(0, 0, 0)
    pf = style.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = pf.space_after = Pt(0)
doc.styles["Title"].font.bold = True
doc.styles["Heading 1"].font.bold = True
doc.styles["Heading 2"].font.bold = True
doc.styles["Caption"].font.bold = False
doc.styles["Caption"].font.italic = True
for style in doc.styles:
    for border in list(style.element.xpath(".//w:pBdr")):
        border.getparent().remove(border)
normal = doc.styles["Normal"].paragraph_format
normal.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
doc.core_properties.title = "Análisis multivariado del tráfico de red UNSW NB15"
doc.core_properties.author = "Mateo Camelino; Ignacio Marinelli; Maximiliano Gimenez; Lisandro Monti"
doc.core_properties.subject = "Proyecto integrador Parte 2 Etapas 0 a 6"
doc.core_properties.language = "es-AR"
settings = doc.settings.element
upd = OxmlElement("w:updateFields"); upd.set(qn("w:val"), "true"); settings.append(upd)
foot = sec.footer.paragraphs[0]
foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
foot._p.append(fld)

table_counter = 0
figure_counter = 0
def p(text="", bold=False, align=None):
    para = doc.add_paragraph()
    para.add_run(str(text)).bold = bold
    if align is not None: para.alignment = align
    return para
def h(text): return doc.add_paragraph(text, "Heading 1")
def sub(text): return doc.add_paragraph(text, "Heading 2")
def pagina(titulo):
    para = h(titulo)
    para.paragraph_format.page_break_before = True
def tabla(headers, rows, title, widths=None):
    global table_counter
    table_counter += 1
    p(f"Tabla {table_counter}. {title}")
    tab = doc.add_table(rows=1, cols=len(headers))
    tab.autofit = False
    widths = widths or [16 / len(headers)] * len(headers)
    for cell, text, w in zip(tab.rows[0].cells, headers, widths):
        cell.width = Cm(w); cell.text = str(text)
    for row in rows:
        cells = tab.add_row().cells
        for j, (cell, value) in enumerate(zip(cells, row)):
            cell.width = Cm(widths[j]); cell.text = str(value)
    for i, row in enumerate(tab.rows):
        pr = row._tr.get_or_add_trPr()
        avoid = OxmlElement("w:cantSplit"); pr.append(avoid)
        if i == 0:
            rep = OxmlElement("w:tblHeader"); pr.append(rep)
        for j, cell in enumerate(row.cells):
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side in ["top", "left", "bottom", "right"]:
                el = OxmlElement("w:" + side); el.set(qn("w:w"), "10" if side in ("top", "bottom") else "40"); el.set(qn("w:type"), "dxa"); margins.append(el)
            tcpr.append(margins)
            borders = OxmlElement("w:tcBorders")
            for side in ["top", "left", "bottom", "right"]:
                el = OxmlElement("w:" + side); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:color"), "D9D9D9"); borders.append(el)
            tcpr.append(borders)
            if i == 0:
                shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "DCE8EF"); tcpr.append(shade)
            for par in cell.paragraphs:
                par.paragraph_format.line_spacing = 1.5
                par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in par.runs:
                    run.font.name = "Calibri"; run.font.size = Pt(11); run.bold = i == 0
    return tab
def figura(file, caption, height=None):
    global figure_counter
    figure_counter += 1
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.line_spacing = 1
    par.paragraph_format.keep_with_next = True
    run = par.add_run()
    shape = run.add_picture(str(F / file), width=Cm(16))
    if height is not None and shape.height > Cm(height):
        shape.width = int(shape.width * Cm(height) / shape.height)
        shape.height = Cm(height)
    shape._inline.docPr.set("descr", caption)
    shape._inline.docPr.set("title", f"Figura {figure_counter}")
    b = OxmlElement("w:pBdr")
    for side in ["top", "left", "bottom", "right"]:
        el = OxmlElement("w:" + side); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:color"), "D9D9D9"); el.set(qn("w:space"), "3"); b.append(el)
    par._p.get_or_add_pPr().append(b)
    cp = doc.add_paragraph(f"Figura {figure_counter}. {caption}", "Caption")
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.keep_with_next = False

# PÁGINA 1. Portada según el PDF de referencia, con los datos corregidos por el usuario.
for _ in range(2): p()
p("UNIVERSIDAD CATÓLICA DE CÓRDOBA", align=WD_ALIGN_PARAGRAPH.CENTER).runs[0].font.size = Pt(16)
p()
p("Facultad de Ingeniería", align=WD_ALIGN_PARAGRAPH.CENTER).runs[0].font.size = Pt(14)
p("Ingeniería en Informática", align=WD_ALIGN_PARAGRAPH.CENTER).runs[0].italic = True
logo = doc.add_paragraph(); logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
logo_shape = logo.add_run().add_picture(str(ROOT / "recursos/escudo_ucc.png"), height=Cm(2.6))
logo_shape._inline.docPr.set("descr", "Escudo de la Universidad Católica de Córdoba, extraído de la portada de referencia")
for _ in range(5): p()
title = doc.add_paragraph("Análisis multivariado del tráfico de red UNSW NB15", "Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p("Proyecto integrador Parte 2", align=WD_ALIGN_PARAGRAPH.CENTER)
for _ in range(3): p()
p("Cátedra: Estadística Multivariada")
p("Docente: Silvina Malpassi")
p("Alumnos:")
for name in ["Mateo Camelino", "Ignacio Marinelli", "Maximiliano Gimenez", "Lisandro Monti"]:
    p("• " + name)
p("Sede Río Cuarto · Segundo semestre de 2026")
p("Fecha de entrega: ____ / ____ / 2026")

# PÁGINA 2. Se completa al final con los epígrafes reales.
pagina("Índices")
index_anchor = doc.paragraphs[-1]
indice_inicio = len(doc.paragraphs)
sub("Secciones")
for text in ["Datos y pregunta del proyecto ........ 3", "Calidad descriptiva y supuestos ........ 4",
    "ACP resultados e interpretación ........ 5–6", "Análisis factorial ........ 7",
    "Conglomerados y perfiles ........ 8–10", "ACS y construcción de X_cat ........ 11",
    "ACM e interpretación ........ 12–13", "Integración y conclusiones ........ 14", "Referencias ........ 15"]:
    p(text)
sub("Tablas")
for text in ["1 Diccionario seleccionado (3); 2 Descriptiva y 3 MSA (4)",
    "4 Autovalores (5); 5 Correlaciones contribución y cos²; 6 Extremos (6)",
    "7 Patrón factorial comunalidades y unicidades (7)",
    "8 Codo y silhouette (8); 9 Ward–k-means; 10 k-means–ACP; 11 Contexto (9)",
    "12 Perfiles originales (10); 13 X_cat y 14 Contingencia ACS (11)",
    "15 Inercias ACM (12); 16 Contribución y cos²; 17 eta² (13)",
    "18 Integración por técnica (14)"]:
    p(text)
sub("Gráficos")
for text in ["1 Datos y correlaciones (4); 2 Codo y biplot ACP (5)",
    "3 Análisis paralelo AF (7); 4 Elección de k y 5 Dendrograma Ward (8)",
    "6 Perfiles y dendrograma de variables (10); 7 Mapa ACS (11)",
    "8 Inercia y 9 Categorías ACM (12); 10 Cluster suplementario y eta² (13)"]:
    p(text)
p("Código completo, tablas ampliadas e imágenes: carpeta del repositorio. El informe comprende las etapas 0 a 6; la regresión logística corresponde a una parte posterior.")

# PÁGINA 3. Datos e hipótesis.
pagina("Datos y pregunta del proyecto")
p("Analizamos patrones y grupos de tráfico para preparar una futura estimación de la probabilidad de ataque. UNSW-NB15 combina comunicaciones normales y ataques sintéticos de laboratorio (Moustafa y Slay, 2015; UNSW, 2021). Cada fila es un flujo. El Excel de trabajo contiene 5.000 filas y 68 columnas; 2.408 son normales y 2.592 ataques. La muestra recibida no se vuelve a muestrear.")
p("Pregunta guía: ¿qué características del tráfico permiten distinguir flujos normales de ataques y estimar su probabilidad mediante un modelo logístico?", bold=True)
p("H1: volumen, temporización y repetición se relacionan dentro de cada bloque y representan dimensiones diferentes. H2: existen grupos de comunicación con diferentes proporciones de ataques. H3: el tipo de tráfico se asocia con protocolo, servicio y estado. Son hipótesis exploratorias formuladas con conocimiento previo de la muestra.")
p("Importamos muestra y diccionario con readxl: 47 columnas numéricas y 21 categóricas según su significado. El diccionario completo de 68 campos se entrega en CSV. La tabla resume X_num; se usan sus versiones log_ = log(x + 1), todas numéricas.")
descripciones = ["Duración del flujo", "Paquetes del origen", "Bytes del origen", "Bytes del destino", "Intervalo entre paquetes", "Variación del intervalo", "Tiempo SYN a SYN-ACK", "Mismo servicio y origen", "Mismo servicio y destino", "Mismo origen", "Mismo origen y puerto destino", "Mismo destino y puerto origen"]
unidades = ["s", "paquetes", "bytes", "bytes", "ms", "ms", "s", "conexiones", "conexiones", "conexiones", "conexiones", "conexiones"]
variables = leer("conjuntos_congelados").query('conjunto == "X_num"').variable.tolist()
tabla(["Nombre sin prefijo log_", "Descripción", "Unidad original"],
    [[abreviar(v), d, u] for v, d, u in zip(variables, descripciones, unidades)],
    "Diccionario de las mediciones seleccionadas", [4.4, 8.5, 3.1])
p("Protocolos, servicios, estados, códigos e identificadores son categóricos; sin unidad significa que no miden una magnitud física. dpkts cuenta paquetes del destino; sbytes y dbytes miden bytes, no paquetes. Las columnas excluidas siguen en el Excel.")
p("Frecuencias: tcp 70,82%, servicio ninguno 60,04% y estado FIN 69,34%. Las frecuencias completas de las categóricas se guardan en frecuencias_categoricas.csv.")

# PÁGINA 4. Calidad y piloto.
pagina("Calidad descriptiva y supuestos")
p("No hay NA, vacíos, duplicados, ID repetidos, negativos, no finitos ni conteos no enteros. is_ftp_login = 4 (id 97597) contradice su definición binaria: se conserva y excluye de X_num/X_cat. No se imputan ceros ni service = ninguno.")
p("Atípicos: bisagras de Tukey Q1 y Q3, fuera de [Q1 − 1,5 RIC; Q3 + 1,5 RIC] en las 12 log_. Se marcan 1.715 flujos (34,30%) y se conservan todos. No se sobrescribe n_atipicos original (1.607 flujos). La revisión y las decisiones se guardan en tablas separadas.")
sumario = leer("resumen_numerico").set_index("variable")
tabla(["Variable", "Media", "Mediana", "Máximo"], [[v, f(sumario.loc[v,'media'],3), f(sumario.loc[v,'mediana'],3), f(sumario.loc[v,'maximo'],3)] for v in ["dur", "sbytes", "dbytes"]], "Mediciones originales seleccionadas; resumen completo en CSV", [3.4, 4.2, 4.2, 4.2])
figura("07_datos_resumen.png", "Duración cruda y logarítmica, boxplots de volumen y correlaciones de X_num.", height=7.2)
piloto = leer("piloto_supuestos").iloc[0]
p(f"Piloto en psych: determinante {f(piloto.determinante,8)} > 0,00001; KMO = {f(piloto.kmo_global,3)}; Bartlett χ²(66) = {f(piloto.bartlett_chi2)}, p < 0,001. No se elimina ninguna variable. Estos controles no prueban normalidad ni independencia.")
msa = leer("msa_por_variable")
tabla(["Variable", "MSA", "Variable", "MSA"], [[abreviar(msa.iloc[i].variable), f(msa.iloc[i].msa,3), abreviar(msa.iloc[i+6].variable), f(msa.iloc[i+6].msa,3)] for i in range(6)], "MSA por variable", [5.7, 2.3, 5.7, 2.3])

# PÁGINA 5. ACP.
pagina("ACP y cantidad de componentes")
p("Aplicamos prcomp(X_num, scale. = TRUE). Aunque todas son log_, sus varianzas difieren: sjit alcanza 13,611 y synack 0,001794. Sin estandarizar CP1 explica 66,43%; estandarizando, 43,29%. La estandarización evita privilegiar una escala de dispersión.")
ev = leer("acp_autovalores")
tabla(["CP", "Autovalor", "%", "% acumulado"], [[int(r.componente), f(r.autovalor,3), f(r.porcentaje), f(r.acumulado)] for r in ev.itertuples()], "Autovalores y varianza del ACP", [2.2, 4.6, 4.6, 4.6])
p("Retenemos tres componentes: son las únicas con autovalor > 1 y explican 74,51% de la varianza. El codo y el análisis paralelo de componentes apoyan esa reducción. El plano CP1–CP2 representa 60,17%; no resume toda la estructura.")
figura("08_acp_codo_biplot.png", "Codo y biplot ACP. Colores por tipo de tráfico y tres extremos señalados.", height=6.2)

# PÁGINA 6. Interpretación ACP.
pagina("Interpretación y extremos del ACP")
p("CP1 se denomina complejidad temporal y volumen frente a repetición: correlaciona positivamente con sjit, sinpkt y volumen, y negativamente con ct_*. CP2 se denomina actividad y volumen compartidos: paquetes, bytes y parte del contexto tienen asociaciones positivas. CP3 opone temporización y repetición por puertos al volumen. Los ejes son combinaciones mixtas; no equivalen a un bloque puro cada uno.")
rc = leer("acp_correlaciones").set_index("variable")
ct = leer("acp_contribuciones").set_index("variable")
cs = leer("acp_cos2").set_index("variable")
tabla(["Variable", "r CP1", "r CP2", "r CP3", "Aporte mayor", "cos² 3 CP"],
    [[abreviar(v), f(rc.loc[v,'PC1'],2), f(rc.loc[v,'PC2'],2), f(rc.loc[v,'PC3'],2),
      f"CP{np.argmax(ct.loc[v].to_numpy())+1}: {f(ct.loc[v].max(),1)}%", f(cs.loc[v].sum(),2)] for v in variables],
    "Correlaciones contribución máxima y calidad de representación", [4.3, 1.9, 1.9, 1.9, 3.3, 2.7])
p("La contribución se calcula como 100 × carga propia²; se compara con 8,33% (100/12). El cos² es la correlación variable–componente al cuadrado. Las tablas completas contienen ambos indicadores por eje. Un cos² bajo exige cautela aunque el punto parezca extremo en el plano.")
ext = leer("acp_extremos")
base = pd.read_excel(ROOT / "muestra_estratificada-simulado-de-ataques.xlsx", sheet_name="muestra").set_index("id")
tabla(["id", "dur (µs)", "spkts", "dpkts", "ct_srv_src", "ct_dst_sport"],
    [[int(r.id), f(r.dur*1e6,0), int(r.spkts), int(r.dpkts), int(base.loc[r.id,'ct_srv_src']), int(base.loc[r.id,'ct_dst_sport_ltm'])] for r in ext.itertuples()],
    "Tres extremos según distancia en las tres CP retenidas", [3.2, 2.7, 1.8, 1.8, 3.2, 3.3])
p("Los tres corresponden a Otros ataques: flujos muy breves, dos paquetes del origen, 114 bytes y ninguna respuesta. Sus conteos de contexto elevados explican su posición negativa en CP1 y positiva en CP2/CP3. La selección usa distancia en las tres CP divididas por sus desviaciones, no solo la apariencia del biplot.")

# PÁGINA 7. AF.
pagina("Análisis factorial y comparación con ACP")
p("Se verifican los supuestos de la página 4. Kaiser en la matriz de correlaciones sugiere 3 ejes, pero el análisis paralelo de factores (100 repeticiones) sugiere 4. Se retienen cuatro por ese criterio y la separación interpretable de dos bloques de contexto.")
figura("09_af_paralelo.png", "Análisis paralelo: los primeros cuatro factores superan la referencia aleatoria.", height=3.0)
p("Minres produce Heywood: comunalidad máxima 1,029 y unicidad −0,029. Se informa el diagnóstico y se contrasta con máxima verosimilitud, varimax y unicidad mínima 0,005. La solución resultante es exploratoria: spkts queda en el límite y el ajuste exacto se rechaza (χ²(24) = 3.653,87; p < 0,001).")
pat = leer("af_patron")
tabla(["Variable", "F1", "F2", "F3", "F4", "h²", "u²"],
    [[abreviar(r.variable), *[f(getattr(r,k),2) for k in ['F1','F2','F3','F4']], f(r.comunalidad,3), f(r.unicidad,3)] for r in pat.itertuples()],
    "Patrón varimax con cargas menores a 0,40 ocultas", [5.2, 1.8, 1.8, 1.8, 1.8, 1.8, 1.8])
p("F1: temporización; F2: volumen; F3: repetición por origen/puertos; F4: repetición por servicio. Explican 22,05%, 21,47%, 16,55% y 13,58% de la varianza total (73,65%). No hay dos cargas ≥ 0,40 en una variable; la menor diferencia entre cargas dominantes es 0,200, apenas sobre el umbral. Los puntajes por regresión se guardan y son estimaciones.")
p("ACP usa varianza total (3 CP); AF, varianza común (4 factores). AF divide el contexto y presenta fragilidad. Preferimos ACP para resumir y agrupar; la interpretación latente requiere cautela.")

# PÁGINA 8. Clustering.
pagina("Conglomerados y elección de k")
p("Calculamos distancias euclidianas sobre las 12 log_ estandarizadas, Ward.D2 y k-means (nstart = 25, semilla 42). Se evaluaron k = 2 a 8. WSS mide dispersión interna; la silhouette se calculó con los 5.000 flujos, sin submuestrear.")
met = leer("cluster_metricas")
tabla(["k", "WSS Ward", "WSS KM", "Sil. Ward", "Sil. KM"],
    [[int(r.k),f(r.wss_ward,0),f(r.wss_kmeans,0),f(r.silhouette_ward,3),f(r.silhouette_kmeans,3)] for r in met.itertuples()],
    "Comparación de candidatos", [1.2,3.7,3.7,3.7,3.7])
figura("10_cluster_seleccion.png", "Codo y silhouette para Ward y k-means en los candidatos k = 2 a 8.", height=4.6)
p("Se elige k = 6 por la mayor silhouette de k-means (0,333), descenso de WSS y perfiles interpretables. La mejora sobre k = 2 (0,328) es solo 0,005: no es un óptimo indiscutible. Ward prefiere k = 2 por silhouette; conservamos esa discrepancia como sensibilidad. Los grupos tienen separación moderada y uno reúne 2,94% de la muestra.")
figura("11_cluster_dendrograma.png", "Ward.D2 truncado a 15 nodos agregados; la línea indica el corte en seis grupos.", height=4.1)

# PÁGINA 9. Comparaciones.
pagina("Comparación de particiones y contexto")
p("Se comparan particiones con tablas cruzadas y Rand ajustado (ARI), independiente de las etiquetas. Ward y k-means alcanzan ARI = 0,813: acuerdo amplio con diferencias en los bordes.")
cr = leer("cluster_cruce_ward_kmeans")
tabla(["Ward / KM", *[f"C{i}" for i in range(1,7)]], [[f"W{i+1}", *[int(x) for x in row]] for i,row in enumerate(cr.to_numpy())],
    "Ward frente a k-means sobre las 12 variables", [3.4,2.1,2.1,2.1,2.1,2.1,2.1])
p("Se repitieron ambos métodos sobre las tres CP sin reestandarizar puntajes: ARI de k-means = 0,708 y Ward = 0,605. Descartar 25,49% de varianza cambia distancias y fronteras.")
cr = leer("cluster_cruce_acp")
tabla(["KM log / ACP", *[f"A{i}" for i in range(1,7)]], [[f"C{i+1}", *[int(x) for x in row]] for i,row in enumerate(cr.to_numpy())],
    "K-means sobre log_ frente a puntajes ACP", [3.4,2.1,2.1,2.1,2.1,2.1,2.1])
contexto = leer("cluster_contexto").pivot(index="cluster", columns="tipo_trafico", values="Freq")
tabla(["Cluster", "Normal", "Expl.", "Fuzz.", "Recon.", "Otros"],
    [[f"C{i}", *[int(contexto.loc[i,c]) for c in ["Normal","Exploits","Fuzzers","Reconnaissance","Otros ataques"]]] for i in range(1,7)],
    "Cruce de k-means con tipo de tráfico", [2.7,2.7,2.7,2.7,2.7,2.5])
p("El contexto se incorpora después de agrupar; no intervino en distancias o centroides. Esta tabla no evalúa predicción.")

# PÁGINA 10. Perfiles.
pagina("Perfiles de grupos y bloques de variables")
pr = leer("cluster_medias_originales").set_index("cluster")
sizes = leer("cluster_tamanos").set_index("cluster")
att = leer("cluster_proporcion_ataques").set_index("cluster")
tabla(["C", "n", "dur s", "spkts", "sbytes", "ct_srv_src", "% ataque"],
    [[i, int(sizes.loc[i,'n']),f(pr.loc[i,'dur'],3),f(pr.loc[i,'spkts'],1),f(pr.loc[i,'sbytes'],0),f(pr.loc[i,'ct_srv_src'],1),f(100*att.loc[i,'x'],1)] for i in range(1,7)],
    "Medias en unidades originales y composición", [0.8,1.8,2.5,2.4,3.2,2.8,2.5])
figura("12_cluster_perfiles_variables.png", "Medias z y dendrograma de variables. Se abrevian los nombres de las 12 log_.", height=6.0)
p("C1, intercambio voluminoso con respuesta rápida: 17,96% de los flujos y 1,45% de ataques. C2, comunicaciones prolongadas y de gran volumen: media de 10,37 s y 72,59% de ataques. C3, repetición extrema sin respuesta: dos paquetes, dbytes = 0 y 100% de ataques; es el grupo pequeño (147 flujos).")
p("C4, temporización elevada con volumen intermedio: 42,62% de la muestra y 71,05% de ataques. C5, comunicaciones breves de bajo volumen: 22,70% y 46,26% de ataques. C6, repetición de servicios con volumen bajo: 5,68% y 34,86% de ataques. Son perfiles descriptivos, no diagnósticos causales.")
p("Ward de variables identifica tres bloques: volumen (spkts, sbytes, dbytes), temporización (dur, sinpkt, sjit, synack) y contexto (cinco ct_*). Coinciden con H1 y con los bloques del AF, aunque este divide el contexto. Las CP combinan esos bloques en ejes ortogonales, en vez de reproducir uno por uno sus grupos.")

# PÁGINA 11. ACS y Xcat.
pagina("Construcción de X_cat y correspondencia simple")
p("X_cat tiene siete factores activos. Las tres discretizaciones usan presencia de respuesta y terciles sin consultar la etiqueta. Se congelan los cortes exactos en el RDS. Ninguna modalidad activa queda por debajo de 5%; el cluster suplementario conserva sus grupos, incluido C3 (2,94%).")
tabla(["Variable", "Modalidades finales"], [
    ["tipo_trafico", "Normal, Exploits, Fuzzers, Reconnaissance, Otros ataques"],
    ["proto_g", "tcp, udp, otros"], ["service_g", "ninguno, http, dns, otros"], ["state_g", "FIN, INT, CON_otros"],
    ["respuesta_destino", "Sin respuesta (dpkts = 0), con respuesta (> 0)"],
    ["duracion_banda", "Corta ≤ 0,018541 s; media ≤ 0,709029333 s; larga > segundo corte"],
    ["bytes_origen_banda", "Bajos ≤ 564; medios ≤ 1.580; altos > 1.580 bytes"]],
    "Variables categóricas y cortes", [4.9,11.1])
p("Para el ACS se cruzaron tipo_trafico y state_g. El chi-cuadrado es 1.637,20 (8 gl), p < 0,001; V de Cramér = 0,405. La frecuencia esperada mínima es 43,28. La asociación es apreciable; el valor p es aproximado por la dependencia potencial entre flujos.")
acs = leer("acs_contingencia")
tabla(["Tipo", "CON_otros", "FIN", "INT"], [[r.tipo_trafico,int(r.CON_otros),int(r.FIN),int(r.INT)] for r in acs.itertuples()],
    "Contingencia para el ACS", [7,3,3,3])
figura("13_acs_mapa.png", "Mapa simétrico ACS, interpretado junto con los porcentajes de fila.", height=4.3)
p("Normal presenta CON_otros en 24,04%, frente a 0,27–3,84% en ataques. Otros ataques presenta INT en 62,30%; Exploits, FIN en 90,45%. La proximidad entre variables no implica distancia directa ni causalidad.")

# PÁGINA 12. ACM dimensiones.
pagina("ACM e inercia de las dimensiones")
p("Se ajustó MCA de FactoMineR con el cluster como suplementario, y ca::mjca(lambda = adjusted) para inercia ajustada. Hay Q = 7 variables, J = 23 modalidades e inercia cruda J/Q − 1 = 2,286. El umbral 1/Q = 0,143 deja seis ejes candidatos; se interpreta el plano de dos por codo y claridad.")
e = leer("acm_autovalores")
tabla(["Dim", "λ crudo", "% crudo", "% ajustado", "% aj. acum."],
    [[int(r.dimension),f(r.autovalor,3),f(r.porcentaje_crudo),f(r.porcentaje_ajustado),f(r.acumulado_ajustado)] for r in e.iloc[:6].itertuples()],
    "Inercia cruda y ajustada del ACM", [1.3,3.4,3.4,3.9,4])
figura("14_acm_inercia.png", "Codo e inercia ajustada, diferentes de la inercia cruda.", height=3.1)
figura("15_acm_categorias.png", "Mapa activo ACM. Tr: tráfico; Pr: protocolo; Sv: servicio; Es: estado; Rp: respuesta; Du: duración; By: bytes.", height=7.8)
p("El plano explica 41,45% crudo y 87,40% ajustado. mjca usa la estructura fuera de la diagonal de Burt: los porcentajes positivos suman 89,90%, sin renormalizarlos a 100. La geometría corresponde al ACM indicador.")

# PÁGINA 13. Interpretación ACM.
pagina("Interpretación del ACM y vínculo con clusters")
cont = leer("acm_categorias_contribuciones").set_index("categoria")
cos = leer("acm_categorias_cos2").set_index("categoria")
rows=[]
for dim in ["Dim.1", "Dim.2"]:
    for cat in cont[dim].nlargest(3).index:
        rows.append([dim[-1],cat.replace("respuesta_destino:","resp:").replace("duracion_banda:","dur:"),f(cont.loc[cat,dim]),f(cos.loc[cat,dim],3)])
tabla(["Dim", "Categoría", "Aporte %", "cos²"], rows,
    "Tres contribuciones principales por dimensión", [1.2,9.2,2.8,2.8])
p("Dim 1, respuesta y complejidad, opone TCP/FIN, respuesta y duración media-larga a comunicaciones cortas e incompletas. Dim 2, estados y servicios, contrapone CON_otros y dns/udp a INT y otros protocolos. Se consideran contribución y cos².")
eta=leer("acm_eta2")
tabla(["Variable", "eta² D1", "eta² D2"], [[r.variable,f(getattr(r,'_1'),3),f(getattr(r,'_2'),3)] for r in eta.iloc[:,:3].itertuples(index=False)],
    "Discriminación de las variables", [9.6,3.2,3.2])
figura("16_acm_clusters_eta.png", "Clusters suplementarios y eta². El cluster no participó en el ajuste del ACM.", height=5.1)
p("C3 es extremo positivo en D1; C5 combina D1 positiva y D2 negativa. C2/C4 se superponen: no hay separación completa. Asociaciones: Otros ataques–sin respuesta (63,18%; Normal 6,89%), Exploits–FIN (90,45%) y Normal–CON_otros (24,04%). Agrupar clases y relacionar variables dependientes puede reforzar estos patrones descriptivos.")

# PÁGINA 14. Integración.
pagina("Integración conclusiones y exposición")
tabla(["Técnica", "Estructura encontrada", "Coincidencias y diferencias"], [
    ["ACP", "3 CP; 74,51% de varianza total", "Combina volumen, tiempos y contexto en ejes ortogonales."],
    ["AF", "4 factores de varianza común", "Reconoce tiempo y volumen; divide contexto. Heywood y límite exigen cautela."],
    ["Clustering", "3 bloques de variables y 6 perfiles de flujos", "ARI Ward/KM 0,813; silhouette 0,333. AF reconoce bloques; las CP cambian fronteras."],
    ["ACS", "Tipo de tráfico × estado; V = 0,405", "Normal–CON_otros y Exploits–FIN coinciden con asociaciones del ACM."],
    ["ACM", "2 ejes interpretados; 87,40% ajustada", "Respuesta, protocolo y estado explican parte de los perfiles; C2/C4 se superponen."]],
    "Comparación integrada", [3.3,5.2,7.5])
p("H1 recibe apoyo descriptivo: se recuperan tres bloques, aunque las CP los mezclan. H2 recibe apoyo parcial: perfiles y proporciones distintos, con separación moderada y poca ventaja de k = 6 sobre k = 2. H3 recibe apoyo: V de Cramér es 0,302 para protocolo, 0,320 para servicio y 0,405 para estado (p < 0,001).")
p("La pregunta guía se responde parcialmente: volumen, tiempos, repetición y respuesta distinguen comportamientos en esta muestra. No se estimó una probabilidad mediante logística en la Parte 2. Esa etapa necesitará separar entrenamiento/prueba antes de aprender transformaciones y cortes, evitar fuga de la etiqueta y excluir TTL por artefactos de laboratorio.")
p("Limitaciones: laboratorio, deduplicación previa, prevalencias no operativas, ventanas compartidas, asimetrías y extremos conservados. AF está en un límite; k depende del criterio; discretizar pierde información y el plano ACM no representa toda la inercia. Las hipótesis exploradas no quedan confirmadas fuera de esta muestra.")
p("Próximos pasos: validar la estabilidad de clusters y cortes, repetir AF con estimación robusta y evaluar una logística con validación independiente y calibración. No inferir causalidad a partir de mapas o clusters.")
p("Exposición propuesta de 18 minutos: Mateo 0–4 (datos/ACP), Ignacio 4–8 (AF y comparación), Maximiliano 8–13 (clusters), Lisandro 13–18 (ACS/ACM y conclusiones). Reservar hasta 2 minutos de cierre sin superar 20. El guion ampliado y el código se entregan separados.")

# PÁGINA 15. Bibliografía.
pagina("Referencias")
refs = [
"Lê, S., Josse, J. y Husson, F. (2008). FactoMineR: An R package for multivariate analysis. Journal of Statistical Software, 25(1). https://doi.org/10.18637/jss.v025.i01",
"Moustafa, N., Creech, G. y Slay, J. (2017). Big data analytics for intrusion detection system: Statistical decision-making using finite Dirichlet mixture models. En Data Analytics and Decision Support for Cybersecurity, 127–156. https://doi.org/10.1007/978-3-319-59439-2_5",
"Moustafa, N. y Slay, J. (2015). UNSW-NB15: A comprehensive data set for network intrusion detection systems. MilCIS, IEEE. https://doi.org/10.1109/MilCIS.2015.7348942",
"Moustafa, N. y Slay, J. (2016). The evaluation of Network Anomaly Detection Systems: Statistical analysis of the UNSW-NB15 dataset and the comparison with the KDD99 dataset. Information Security Journal, 25(1–3), 18–31. https://doi.org/10.1080/19393555.2015.1125974",
"Moustafa, N., Slay, J. y Creech, G. (2017). Novel geometric area analysis technique for anomaly detection using trapezoidal area estimation on large-scale networks. IEEE Transactions on Big Data. https://doi.org/10.1109/TBDATA.2017.2715166",
"Nenadic, O. y Greenacre, M. (2007). Correspondence analysis in R, with two- and three-dimensional graphics: The ca package. Journal of Statistical Software, 20(3). https://doi.org/10.18637/jss.v020.i03",
"Revelle, W. (2026). psych: Procedures for psychological, psychometric, and personality research (versión 2.6.5). https://personality-project.org/r/psych/",
"Sarhan, M., Layeghy, S., Moustafa, N. y Portmann, M. (2021). NetFlow datasets for machine learning-based network intrusion detection systems. Big Data Technologies and Applications, 117–135. https://doi.org/10.1007/978-3-030-72802-1_9",
"UNSW (2021). The UNSW-NB15 Dataset. https://research.unsw.edu.au/projects/unsw-nb15-dataset",
"Universidad Católica de Córdoba (2026). Proyecto integrador Parte 2. Consigna de Estadística Multivariada, sede Río Cuarto. Material de cátedra proporcionado por el grupo."
]
for ref in refs:
    par=p(ref); par.paragraph_format.first_line_indent = Cm(-.6); par.paragraph_format.left_indent=Cm(.6)
p("Software: R 4.6.1; psych 2.6.5; FactoMineR 2.17; ca 0.72; cluster 2.1.8.2. Gráficos: Matplotlib. Semilla: 42. Ejecución: source(\"R/ejecutar_todo.R\", encoding = \"UTF-8\"); luego python/graficos_parte2.py y python/crear_informe.py.")

doc.save(OUT)
print(OUT)
print(f"Páginas planificadas: 15. Tablas: {table_counter}. Figuras: {figure_counter}. Verificar render final.")
